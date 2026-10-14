import streamlit as st
import pandas as pd
import os
import re
from datetime import datetime, timedelta
import calendar
import glob
import subprocess
import io
import json

# openpyxlのスタイリング設定用モジュール
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

# ページ設定
st.set_page_config(
    page_title="売上月報・店舗管理システム",
    page_icon="📊",
    layout="wide"
)

OUTPUT_DIR = os.path.abspath("./output")
DATA_DIR = os.path.abspath("./data")
os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(DATA_DIR, exist_ok=True)

BUDGET_PREV_PATH = os.path.join(DATA_DIR, "budget_prev.csv")
DELIVERY_DATA_PATH = os.path.join(DATA_DIR, "delivery_sales.csv")
CREDENTIALS_PATH = os.path.join(DATA_DIR, "credentials.json")

STORES = [
    "三田おにぎり 本店",
    "おにぎり維心 恵比寿店",
    "おにぎり維心 五反田店",
    "おにぎり維心 中目黒店"
]
TIME_SLOTS = ["早番", "中番", "遅番"]

# デフォルトの権限（ロール）定義
DEFAULT_ROLES = {
    "admin": {
        "name": "管理者",
        "password": "admin123",
        "store": "すべて（切替可能）",
        "allowed_modes": [
            "📊 本部用（月報一覧表示）",
            "📱 店舗用（デリバリー入力）",
            "📋 エクセルコピペ用データ生成",
            "📄 報告書作成",
            "⚙️ 設定（権限・パスワード管理）"
        ],
        "months_limit": "all"
    },
    "hombu": {
        "name": "本部",
        "password": "hombu123",
        "store": "すべて（切替可能）",
        "allowed_modes": [
            "📊 本部用（月報一覧表示）",
            "📱 店舗用（デリバリー入力）",
            "📋 エクセルコピペ用データ生成",
            "📄 報告書作成"
        ],
        "months_limit": "all"
    },
    "mita_tenpo": {
        "name": "三田店店舗",
        "password": "mita123",
        "store": "三田おにぎり 本店",
        "allowed_modes": [
            "📊 本部用（月報一覧表示）",
            "📱 店舗用（デリバリー入力）"
        ],
        "months_limit": 3
    }
}

# 認証情報の読み込み・保存関数（自動マイグレーション付き）
def load_credentials():
    if os.path.exists(CREDENTIALS_PATH):
        try:
            with open(CREDENTIALS_PATH, "r", encoding="utf-8") as f:
                data = json.load(f)
                roles = data.get("roles", DEFAULT_ROLES)
                if "mita_staff" in roles and "mita_tenpo" not in roles:
                    roles["mita_tenpo"] = roles.pop("mita_staff")
                    save_credentials(roles)
                return roles
        except Exception:
            pass
    return DEFAULT_ROLES

def save_credentials(roles):
    data = {"roles": roles}
    with open(CREDENTIALS_PATH, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=4)

ROLES = load_credentials()

def normalize_time_slot(val):
    s = str(val)
    if "早" in s: return "早番"
    elif "中" in s: return "中番"
    elif "遅" in s: return "遅番"
    return s

def init_data_files():
    if not os.path.exists(BUDGET_PREV_PATH):
        records = []
        for day in range(1, 32):
            date_str = f"2026-10-{day:02d}"
            for store in STORES:
                for slot in TIME_SLOTS:
                    if slot == "早番":
                        b_sales, p_sales, p_cust = 88000, 77645, 90
                    elif slot == "中番":
                        b_sales, p_sales, p_cust = 157000, 153848, 194
                    else:
                        b_sales, p_sales, p_cust = 55000, 48521, 78

                    records.append({
                        "date": date_str, "store": store, "time_slot": slot,
                        "budget_sales": b_sales, "prev_sales": p_sales, "prev_cust": p_cust
                    })
        pd.DataFrame(records).to_csv(BUDGET_PREV_PATH, index=False, encoding="utf-8-sig")

    if not os.path.exists(DELIVERY_DATA_PATH):
        pd.DataFrame(columns=["date", "store", "time_slot", "uber", "rocket_now", "demaekan", "other"]).to_csv(DELIVERY_DATA_PATH, index=False, encoding="utf-8-sig")

init_data_files()

def load_budget_prev_data():
    df = pd.read_csv(BUDGET_PREV_PATH)
    df["time_slot"] = df["time_slot"].apply(normalize_time_slot)
    return df

def load_delivery_data():
    df = pd.read_csv(DELIVERY_DATA_PATH)
    df["time_slot"] = df["time_slot"].apply(normalize_time_slot)
    if "rn" in df.columns and "rocket_now" not in df.columns:
        df = df.rename(columns={"rn": "rocket_now"})
    return df

def save_delivery_data(date_str, store, time_slot, uber, rocket_now, demaekan, other):
    df = load_delivery_data()
    df = df[~((df["date"] == date_str) & (df["store"] == store) & (df["time_slot"] == time_slot))]
    new_row = pd.DataFrame([{
        "date": date_str, "store": store, "time_slot": time_slot,
        "uber": uber, "rocket_now": rocket_now, "demaekan": demaekan, "other": other
    }])
    df = pd.concat([df, new_row], ignore_index=True)
    df.to_csv(DELIVERY_DATA_PATH, index=False, encoding="utf-8-sig")

def parse_int(val):
    if pd.isna(val): return 0
    clean_str = re.sub(r"[^\d-]", "", str(val))
    try: return int(clean_str)
    except ValueError: return 0

def find_and_read_airregi(date_str):
    d_dash, d_nodash, d_under = date_str, date_str.replace("-", ""), date_str.replace("-", "_")
    all_files = glob.glob(os.path.join(OUTPUT_DIR, "*.csv"))
    matched = [f for f in all_files if d_dash in os.path.basename(f) or d_nodash in os.path.basename(f) or d_under in os.path.basename(f)]
    if not matched: return None, "ファイルなし"
    matched.sort(key=lambda x: os.path.getmtime(x), reverse=True)
    try:
        return pd.read_csv(matched[0], header=None), os.path.basename(matched[0])
    except Exception as e:
        return None, str(e)

def get_short_store_key(store_name):
    for key in ["三田", "恵比寿", "五反田", "中目黒"]:
        if key in store_name: return key
    return store_name

def match_slot_row(store_df, slot_name):
    short_slot = slot_name[0] if slot_name else ""
    return store_df[store_df[1].astype(str).str.contains(short_slot, na=False)]

def generate_month_df(selected_store, year, month, target_slot_keyword, bp_df, deliv_df):
    _, max_days = calendar.monthrange(year, month)
    rows = []
    for day in range(1, max_days + 1):
        date_str = f"{year:04d}-{month:02d}-{day:02d}"
        bp_sub = bp_df[(bp_df["date"] == date_str) & (bp_df["store"] == selected_store)]
        if target_slot_keyword: bp_sub = bp_sub[bp_sub["time_slot"] == target_slot_keyword]

        p_sales = parse_int(bp_sub["prev_sales"].sum()) if not bp_sub.empty else 0
        b_sales = parse_int(bp_sub["budget_sales"].sum()) if not bp_sub.empty else 0
        p_cust = parse_int(bp_sub["prev_cust"].sum()) if not bp_sub.empty else 0

        air_sales, air_cust, air_other_sales = 0, 0, 0
        raw_df, _ = find_and_read_airregi(date_str)
        if raw_df is not None:
            try:
                short_key = get_short_store_key(selected_store)
                store_df = raw_df[raw_df[0].astype(str).str.contains(short_key, na=False)]
                matched = store_df[store_df[1].astype(str).str.strip() == "合計"] if target_slot_keyword is None else match_slot_row(store_df, target_slot_keyword)
                if not matched.empty:
                    air_sales = parse_int(matched.iloc[0, 2])
                    air_cust = parse_int(matched.iloc[0, 5])
                    if matched.shape[1] > 11: air_other_sales = parse_int(matched.iloc[0, 11])
            except Exception: pass

        deliv_sub = deliv_df[(deliv_df["date"] == date_str) & (deliv_df["store"] == selected_store)]
        if target_slot_keyword: deliv_sub = deliv_sub[deliv_sub["time_slot"] == target_slot_keyword]
        
        day_delivery_sales = int(pd.to_numeric(deliv_sub["uber"], errors="coerce").fillna(0).sum() +
                                 pd.to_numeric(deliv_sub["rocket_now"], errors="coerce").fillna(0).sum() +
                                 pd.to_numeric(deliv_sub["demaekan"], errors="coerce").fillna(0).sum()) if not deliv_sub.empty else 0
        deliv_other_sales = int(pd.to_numeric(deliv_sub["other"], errors="coerce").fillna(0).sum()) if not deliv_sub.empty else 0

        total_other_sales = air_other_sales + deliv_other_sales
        actual_sales = air_sales + day_delivery_sales + total_other_sales
        actual_cust = air_cust

        prev_ratio = f"{(actual_sales / p_sales * 100):.1f}%" if p_sales > 0 and actual_sales != 0 else ("-" if actual_sales == 0 else "0.0%")
        budget_ratio = f"{(actual_sales / b_sales * 100):.1f}%" if b_sales > 0 and actual_sales != 0 else ("-" if actual_sales == 0 else "0.0%")
        p_price = int(p_sales / p_cust) if p_cust > 0 else 0
        a_price = int(actual_sales / actual_cust) if actual_cust > 0 else 0
        cust_ratio = f"{(actual_cust / p_cust * 100):.1f}%" if p_cust > 0 and actual_cust > 0 else "-"
        price_ratio = f"{(a_price / p_price * 100):.1f}%" if p_price > 0 and a_price != 0 else "-"
        weekday_kanji = ["月", "火", "水", "木", "金", "土", "日"][pd.to_datetime(date_str).weekday()]

        rows.append({
            "日": day, "曜日": weekday_kanji, "前年売上": p_sales, "予算売上": b_sales,
            "実績売上": actual_sales, "前年比": prev_ratio, "予算比": budget_ratio,
            "店頭売上": air_sales, "デリバリー売上": day_delivery_sales, "その他売上": total_other_sales,
            "前年客数": p_cust, "実績客数": actual_cust, "客数前年比": cust_ratio,
            "前年単価": p_price, "実績単価": a_price, "単価前年比": price_ratio
        })
    return pd.DataFrame(rows)

def calculate_summary_rows(df_slot):
    df_act = df_slot[df_slot["実績売上"] != 0]
    if df_act.empty: df_act = df_slot.iloc[:1]
    num_cols = ["前年売上", "予算売上", "実績売上", "店頭売上", "デリバリー売上", "その他売上", "前年客数", "実績客数"]
    
    def make_row(label, vals, p_ratio, b_ratio, c_ratio, p_price, a_price, pr_ratio):
        return {
            "日": label, "曜日": "", "前年売上": vals["前年売上"], "予算売上": vals["予算売上"], "実績売上": vals["実績売上"],
            "前年比": p_ratio, "予算比": b_ratio, "店頭売上": vals["店頭売上"], "デリバリー売上": vals["デリバリー売上"], "その他売上": vals["その他売上"],
            "前年客数": vals["前年客数"], "実績客数": vals["実績客数"], "客数前年比": c_ratio, "前年単価": p_price, "実績単価": a_price, "単価前年比": pr_ratio
        }

    sum_v = {col: int(pd.to_numeric(df_act[col], errors="coerce").fillna(0).sum()) for col in num_cols}
    s_pp = int(sum_v['前年売上'] / sum_v['前年客数']) if sum_v['前年客数'] > 0 else 0
    s_ap = int(sum_v['実績売上'] / sum_v['実績客数']) if sum_v['実績客数'] > 0 else 0
    row_sum = make_row("合計", sum_v, f"{(sum_v['実績売上']/sum_v['前年売上']*100):.1f}%" if sum_v['前年売上']>0 else "-",
                       f"{(sum_v['実績売上']/sum_v['予算売上']*100):.1f}%" if sum_v['予算売上']>0 else "-",
                       f"{(sum_v['実績客数']/sum_v['前年客数']*100):.1f}%" if sum_v['前年客数']>0 else "-", s_pp, s_ap,
                       f"{(s_ap/s_pp*100):.1f}%" if s_pp>0 else "-")

    ave_v = {col: int(pd.to_numeric(df_act[col], errors="coerce").fillna(0).mean()) for col in num_cols}
    a_pp = int(ave_v['前年売上'] / ave_v['前年客数']) if ave_v['前年客数'] > 0 else 0
    a_ap = int(ave_v['実績売上'] / ave_v['実績客数']) if ave_v['実績客数'] > 0 else 0
    row_ave = make_row("AVE", ave_v, f"{(ave_v['実績売上']/ave_v['前年売上']*100):.1f}%" if ave_v['前年売上']>0 else "-",
                       f"{(ave_v['実績売上']/ave_v['予算売上']*100):.1f}%" if ave_v['予算売上']>0 else "-",
                       f"{(ave_v['実績客数']/ave_v['前年客数']*100):.1f}%" if ave_v['前年客数']>0 else "-", a_pp, a_ap,
                       f"{(a_ap/a_pp*100):.1f}%" if a_pp>0 else "-")

    min_v = {col: int(pd.to_numeric(df_act[col], errors="coerce").fillna(0).min()) for col in num_cols}
    m_pp = int(pd.to_numeric(df_act["前年単価"], errors="coerce").fillna(0).min())
    m_ap = int(pd.to_numeric(df_act["実績単価"], errors="coerce").fillna(0).min())
    row_min = make_row("MIN", min_v, f"{(min_v['実績売上']/min_v['前年売上']*100):.1f}%" if min_v['前年売上']>0 else "-",
                       f"{(min_v['実績売上']/min_v['予算売上']*100):.1f}%" if min_v['予算売上']>0 else "-",
                       f"{(min_v['実績客数']/min_v['前年客数']*100):.1f}%" if min_v['前年客数']>0 else "-", m_pp, m_ap,
                       f"{(m_ap/m_pp*100):.1f}%" if m_pp>0 else "-")

    max_v = {col: int(pd.to_numeric(df_act[col], errors="coerce").fillna(0).max()) for col in num_cols}
    mx_pp = int(pd.to_numeric(df_act["前年単価"], errors="coerce").fillna(0).max())
    mx_ap = int(pd.to_numeric(df_act["実績単価"], errors="coerce").fillna(0).max())
    row_max = make_row("MAX", max_v, f"{(max_v['実績売上']/max_v['前年売上']*100):.1f}%" if max_v['前年売上']>0 else "-",
                       f"{(max_v['実績売上']/max_v['予算売上']*100):.1f}%" if max_v['予算売上']>0 else "-",
                       f"{(max_v['実績客数']/max_v['前年客数']*100):.1f}%" if max_v['前年客数']>0 else "-", mx_pp, mx_ap,
                       f"{(mx_ap/mx_pp*100):.1f}%" if mx_pp>0 else "-")

    return [row_sum, row_ave, row_min, row_max]

def create_multi_sheet_excel(selected_store, year, month, bp_df, deliv_df):
    output = io.BytesIO()
    slots = [("合計", None), ("早番", "早番"), ("中番", "中番"), ("遅番", "遅番")]
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="游ゴシック", size=11, bold=True, color="FFFFFF")
    even_row_fill = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")
    odd_row_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    data_font = Font(name="游ゴシック", size=10)
    summary_label_font = Font(name="游ゴシック", size=10, bold=True)
    center_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"),
                         top=Side(style="thin", color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9"))

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        workbook = writer.book
        for sheet_name, slot_kw in slots:
            df_slot = generate_month_df(selected_store, year, month, slot_kw, bp_df, deliv_df)
            cols = list(df_slot.columns)
            ws = workbook.create_sheet(title=sheet_name)
            ws.cell(row=1, column=1, value=f"{selected_store} {year}年{month}月度（{sheet_name}）")
            ws.cell(row=1, column=1).font = Font(name="游ゴシック", size=13, bold=True, color="1F4E78")
            
            for col_idx, col_name in enumerate(cols, start=1):
                cell = ws.cell(row=3, column=col_idx, value=col_name if col_idx > 2 else "")
                cell.fill = header_fill; cell.font = header_font; cell.alignment = center_alignment; cell.border = thin_border
            
            for s_idx, s_row in enumerate(calculate_summary_rows(df_slot), start=4):
                fill = even_row_fill if s_idx % 2 == 0 else odd_row_fill
                for col_idx, col_name in enumerate(cols, start=1):
                    cell = ws.cell(row=s_idx, column=col_idx, value=s_row.get(col_name, "") if col_idx != 2 else "")
                    cell.font = summary_label_font if col_idx == 1 else data_font
                    cell.fill = fill; cell.alignment = center_alignment; cell.border = thin_border

            for col_idx, col_name in enumerate(cols, start=1):
                cell = ws.cell(row=9, column=col_idx, value=col_name)
                cell.fill = header_fill; cell.font = header_font; cell.alignment = center_alignment; cell.border = thin_border
                
            for r_idx, row_data in enumerate(df_slot.to_dict(orient="records"), start=10):
                fill = even_row_fill if r_idx % 2 == 0 else odd_row_fill
                for col_idx, col_name in enumerate(cols, start=1):
                    cell = ws.cell(row=r_idx, column=col_idx, value=row_data[col_name])
                    cell.fill = fill; cell.font = data_font; cell.alignment = center_alignment; cell.border = thin_border

            ws.column_dimensions['A'].width = 8; ws.column_dimensions['B'].width = 8
            for col in list(ws.columns)[2:]:
                max_len = max(len(str(cell.value or '')) for cell in list(col)[2:])
                ws.column_dimensions[col[0].column_letter].width = max(max_len + 4, 12)

        if "Sheet" in workbook.sheetnames: workbook.remove(workbook["Sheet"])
    return output.getvalue()


# ==========================================
# ログイン認証画面（パスワード入力のみ）
# ==========================================
if "logged_in" not in st.session_state:
    st.session_state["logged_in"] = False
if "current_role_key" not in st.session_state:
    st.session_state["current_role_key"] = None

if not st.session_state["logged_in"]:
    st.title("🔐 ログイン認証")
    st.markdown("パスワードを入力してログインしてください。")
    st.markdown("---")

    input_pass = st.text_input("パスワード", type="password")

    if st.button("ログイン", type="primary"):
        matched_role_key = None
        for r_key, r_info in ROLES.items():
            if input_pass == r_info["password"]:
                matched_role_key = r_key
                break

        if matched_role_key:
            st.session_state["logged_in"] = True
            st.session_state["current_role_key"] = matched_role_key
            st.success("🎉 ログイン成功しました！")
            st.rerun()
        else:
            st.error("❌ パスワードが間違っています。")
    st.stop()


# ==========================================
# ログイン後のメインアプリケーション
# ==========================================
role_key = st.session_state["current_role_key"]
role_info = ROLES.get(role_key, ROLES["admin"])

# 店舗スコープの決定
assigned_store = role_info["store"]
is_all_stores = (assigned_store == "すべて（切替可能）")

if "selected_store" not in st.session_state:
    st.session_state["selected_store"] = STORES[0] if is_all_stores else assigned_store
if "selected_year_month" not in st.session_state:
    st.session_state["selected_year_month"] = "2026-10"
if "target_date_state" not in st.session_state:
    st.session_state["target_date_state"] = datetime.now().date()

if not is_all_stores:
    st.session_state["selected_store"] = assigned_store

# --- サイドバー設定 ---
st.sidebar.title("メニュー")
# ※ 権限の表示部分を削除しました
if not is_all_stores:
    st.sidebar.markdown(f"🏬 担当店舗: **{assigned_store}**")

st.sidebar.markdown("---")
st.sidebar.subheader("店舗・期間設定")

if is_all_stores:
    st.sidebar.selectbox("閲覧店舗", STORES, key="selected_store")
else:
    st.sidebar.info(f"📍 担当店舗: {assigned_store} （固定）")

# 閲覧可能月（months_limit）の制限ロジック
months_limit = role_info.get("months_limit", "all")
if months_limit == "all":
    available_months = ["2026-10", "2026-09", "2026-08", "2026-07"]
else:
    try:
        limit_n = int(months_limit)
    except:
        limit_n = 3
    
    current_dt = datetime(2026, 10, 1)
    available_months = []
    for i in range(limit_n):
        y = current_dt.year
        m = current_dt.month - i
        while m <= 0:
            m += 12
            y -= 1
        available_months.append(f"{y:04d}-{m:02d}")

st.sidebar.selectbox("対象年月", available_months, key="selected_year_month")

st.sidebar.markdown("---")
app_mode = st.sidebar.radio("機能を選択してください", role_info["allowed_modes"])

st.sidebar.markdown("---")
if st.sidebar.button("🚪 ログアウト"):
    st.session_state["logged_in"] = False
    st.session_state["current_role_key"] = None
    st.rerun()

selected_store = st.session_state["selected_store"]
selected_year_month = st.session_state["selected_year_month"]
year, month = map(int, selected_year_month.split("-"))

# ==========================================
# 各画面の描画
# ==========================================
if app_mode == "⚙️ 設定（権限・パスワード管理）":
    st.title("⚙️ 権限・パスワード・閲覧月管理")
    st.markdown("権限（ロール）ごとのパスワード、利用可能機能、月報の閲覧可能月数を自由に追加・編集できます。")
    st.markdown("---")

    updated_roles = {}
    for r_key, r_val in ROLES.items():
        st.markdown(f"**権限キー (ID): `{r_key}`**")
        col_r1, col_r2, col_r3 = st.columns(3)
        with col_r1:
            r_name = st.text_input(f"表示名 ({r_key})", value=r_val["name"], key=f"name_{r_key}")
        with col_r2:
            r_pass = st.text_input(f"パスワード ({r_key})", value=r_val["password"], type="password", key=f"pass_{r_key}")
        with col_r3:
            limit_val = r_val.get("months_limit", "all")
            r_limit = st.text_input(f"閲覧可能月 ('all' または 数字)", value=str(limit_val), key=f"limit_{r_key}")
        
        r_store = st.selectbox(f"担当店舗 ({r_key})", ["すべて（切替可能）"] + STORES, index=0 if r_val["store"]=="すべて（切替可能）" else STORES.index(r_val["store"])+1, key=f"store_{r_key}")

        all_possible_modes = [
            "📊 本部用（月報一覧表示）",
            "📱 店舗用（デリバリー入力）",
            "📋 エクセルコピペ用データ生成",
            "📄 報告書作成",
            "⚙️ 設定（権限・パスワード管理）"
        ]
        r_modes = st.multiselect(f"利用可能機能 ({r_key})", all_possible_modes, default=r_val["allowed_modes"], key=f"modes_{r_key}")
        
        updated_roles[r_key] = {
            "name": r_name,
            "password": r_pass,
            "store": r_store,
            "allowed_modes": r_modes,
            "months_limit": int(r_limit) if r_limit.isdigit() else "all"
        }
        st.markdown("---")

    st.subheader("➕ 新しい権限を追加する")
    new_key = st.text_input("新規権限キー (例: ebisu_tenpo ※半角英数字)", key="new_role_key")
    new_name = st.text_input("新規権限表示名 (例: 恵比寿店店舗)", key="new_role_name")
    new_pass = st.text_input("新規パスワード", type="password", key="new_role_pass")
    new_store = st.selectbox("新規担当店舗", ["すべて（切替可能）"] + STORES, key="new_role_store")
    new_limit = st.text_input("閲覧可能月 ('all' または 数字)", value="3", key="new_role_limit")
    new_modes = st.multiselect("新規利用可能機能", [
        "📊 本部用（月報一覧表示）",
        "📱 店舗用（デリバリー入力）",
        "📋 エクセルコピペ用データ生成",
        "📄 報告書作成"
    ], default=["📊 本部用（月報一覧表示）", "📱 店舗用（デリバリー入力）"], key="new_role_modes")

    if st.button("💾 設定を保存する", type="primary"):
        if new_key and new_key not in updated_roles:
            updated_roles[new_key] = {
                "name": new_name or new_key,
                "password": new_pass or "1234",
                "store": new_store,
                "allowed_modes": new_modes,
                "months_limit": int(new_limit) if new_limit.isdigit() else "all"
            }
        save_credentials(updated_roles)
        global ROLES
        ROLES = updated_roles
        st.success("🎉 権限設定を保存しました！変更が即時反映されました。")
        st.rerun()

elif app_mode == "📱 店舗用（デリバリー入力）":
    st.title("🛵 デリバリー売上 入力フォーム")
    st.markdown(f"現在選択中の店舗: **{selected_store}** | 対象の日付、時間帯を選択してデリバリー売上を入力・更新します。")

    col_target1, col_target2 = st.columns(2)
    with col_target1:
        st.info(f"📍 対象店舗: **{selected_store}**")
    with col_target2:
        selected_slot = st.selectbox("時間帯を選択", TIME_SLOTS)

    selected_date_val = st.date_input("対象日付", value=st.session_state["target_date_state"], key="deliv_date_input")
    if selected_date_val != st.session_state["target_date_state"]:
        st.session_state["target_date_state"] = selected_date_val
    selected_date = selected_date_val.strftime("%Y-%m-%d")
    
    st.markdown("---")
    
    deliv_df = load_delivery_data()
    existing_row = deliv_df[(deliv_df["date"] == selected_date) & (deliv_df["store"] == selected_store) & (deliv_df["time_slot"] == selected_slot)]
    is_already_registered = not existing_row.empty

    default_uber, default_rn, default_demae, default_other = 0, 0, 0, 0
    if is_already_registered:
        if "uber" in existing_row.columns: default_uber = parse_int(existing_row["uber"].values[0])
        if "rocket_now" in existing_row.columns: default_rn = parse_int(existing_row["rocket_now"].values[0])
        if "demaekan" in existing_row.columns: default_demae = parse_int(existing_row["demaekan"].values[0])
        if "other" in existing_row.columns: default_other = parse_int(existing_row["other"].values[0])
        st.info(f"ℹ️ この日時（{selected_date} / {selected_slot}）のデータはすでに登録されています。変更して保存すると上書きされます。")

    slot_key = f"{selected_date}_{selected_store}_{selected_slot}"
    col_d1, col_d2 = st.columns(2)
    with col_d1:
        uber_val = st.number_input("Uber Eats 売上 (円)", value=default_uber, step=500, key=f"input_uber_{slot_key}")
        rn_val = st.number_input("Rocket now 売上 (円)", value=default_rn, step=500, key=f"input_rn_{slot_key}")
    with col_d2:
        demaekan_val = st.number_input("出前館 売上 (円)", value=default_demae, step=500, key=f"input_demae_{slot_key}")
        other_val = st.number_input("その他 売上 (円)", value=default_other, step=500, key=f"input_other_{slot_key}")

    st.markdown("---")
    total_delivery = uber_val + rn_val + demaekan_val + other_val
    if total_delivery < 0:
        st.markdown(f"💡 この時間帯の売上合計: <span style='color: #c0392b; font-weight: bold; font-size: 1.15em;'>¥{total_delivery:,}</span>", unsafe_allow_html=True)
    else:
        st.markdown(f"💡 この時間帯の売上合計: **¥{total_delivery:,}**")

    if st.button("💾 デリバリー売上を保存する", type="primary"):
        save_delivery_data(selected_date, selected_store, selected_slot, uber_val, rn_val, demaekan_val, other_val)
        st.success(f"🎉 {selected_store} ({selected_date} / {selected_slot}) のデリバリー売上を保存しました！")
        st.rerun()

elif app_mode == "📋 エクセルコピペ用データ生成":
    st.title("📋 エクセル実績入力用データ生成")
    st.markdown(f"現在選択中の店舗: **{selected_store}** | エクセルシートへ簡単にコピペするためのデータを生成します。")

    target_date_val = st.date_input("対象日付", value=st.session_state["target_date_state"], key="copy_date_input")
    if target_date_val != st.session_state["target_date_state"]:
        st.session_state["target_date_state"] = target_date_val
    target_date_str = target_date_val.strftime("%Y-%m-%d")

    raw_df, file_status = find_and_read_airregi(target_date_str)
    if raw_df is None:
        st.warning(f"⚠️ **{target_date_str}** のAirレジCSVファイルが見つかりません。（Status: {file_status}）")

    def get_slot_metrics(store_name, slot_name, df_data):
        air_sales, air_cust = 0, 0
        if df_data is not None:
            try:
                short_key = get_short_store_key(store_name)
                store_df = df_data[df_data[0].astype(str).str.contains(short_key, na=False)]
                matched = match_slot_row(store_df, slot_name)
                if not matched.empty:
                    air_sales = parse_int(matched.iloc[0, 2])
                    air_cust = parse_int(matched.iloc[0, 5])
            except Exception: pass
        return air_sales, air_cust

    def get_slot_delivery_total(store_name, date_str, slot_name):
        deliv_df = load_delivery_data()
        sub = deliv_df[(deliv_df["date"].astype(str).str.strip() == date_str.strip()) & 
                       (deliv_df["store"].astype(str).str.strip() == store_name.strip()) & 
                       (deliv_df["time_slot"].astype(str).str.strip() == slot_name.strip())]
        if sub.empty: return 0
        return int(pd.to_numeric(sub["uber"], errors="coerce").fillna(0).sum() +
                   pd.to_numeric(sub["rocket_now"], errors="coerce").fillna(0).sum() +
                   pd.to_numeric(sub["demaekan"], errors="coerce").fillna(0).sum())

    early_d = get_slot_delivery_total(selected_store, target_date_str, "早番")
    mid_d = get_slot_delivery_total(selected_store, target_date_str, "中番")
    late_d = get_slot_delivery_total(selected_store, target_date_str, "遅番")
    
    early_sales, early_cust = get_slot_metrics(selected_store, "早番", raw_df)
    mid_sales, mid_cust = get_slot_metrics(selected_store, "中番", raw_df)
    late_sales, late_cust = get_slot_metrics(selected_store, "遅番", raw_df)

    bp_df_tmp = load_budget_prev_data()
    deliv_df_tmp = load_delivery_data()
    df_month_all = generate_month_df(selected_store, target_date_val.year, target_date_val.month, None, bp_df_tmp, deliv_df_tmp)
    day_row = df_month_all[df_month_all["日"] == target_date_val.day]
    target_other_sales = int(day_row["その他売上"].values[0]) if not day_row.empty else 0

    parts = [str(early_cust), str(early_sales), str(early_d), str(mid_cust), str(mid_sales), str(mid_d), str(late_cust), str(late_sales), str(late_d), str(target_other_sales)]
    tabs_list = ["\t" * 4, "\t" * 1, "\t" * 3, "\t" * 4, "\t" * 1, "\t" * 3, "\t" * 4, "\t" * 1, "\t" * 1]
    tsv_text = "".join([part + tabs_list[i] if i < len(tabs_list) else part for i, part in enumerate(parts)])
    
    st.success(f"🎉 {target_date_str} ({selected_store}) のデータを自動生成しました！")
    st.code(tsv_text, language="text")

elif app_mode == "📄 報告書作成":
    st.title("📄 LINE報告書作成")
    st.markdown("指定した日付（または直近実績日まで）の全店舗の数値を集計し、LINE報告用のテキストを一括生成します。")
    st.markdown("---")

    report_date = st.date_input("集計基準日（この日までを累計します）", value=st.session_state["target_date_state"], key="report_date_input")
    if report_date != st.session_state["target_date_state"]:
        st.session_state["target_date_state"] = report_date
    
    if st.button("📝 報告用テキストを生成する", type="primary"):
        bp_df = load_budget_prev_data()
        deliv_df = load_delivery_data()
        _, max_days = calendar.monthrange(year, month)
        target_end_day = min(report_date.day, max_days)

        report_lines = ["お疲れ様です。", f"{report_date.month}月{report_date.day}日時点の中間報告です。", ""]

        for store_name in STORES:
            total_p, total_b, total_act = 0, 0, 0
            for day in range(1, target_end_day + 1):
                date_str = f"{year:04d}-{month:02d}-{day:02d}"
                bp_sub = bp_df[(bp_df["date"] == date_str) & (bp_df["store"] == store_name)]
                p_sales = parse_int(bp_sub["prev_sales"].sum()) if not bp_sub.empty else 0
                b_sales = parse_int(bp_sub["budget_sales"].sum()) if not bp_sub.empty else 0

                air_sales, air_other_sales = 0, 0
                raw_df, _ = find_and_read_airregi(date_str)
                if raw_df is not None:
                    try:
                        short_key = get_short_store_key(store_name)
                        store_df = raw_df[raw_df[0].astype(str).str.contains(short_key, na=False)]
                        matched = store_df[store_df[1].astype(str).str.strip() == "合計"]
                        if not matched.empty:
                            air_sales = parse_int(matched.iloc[0, 2])
                            if matched.shape[1] > 11: air_other_sales = parse_int(matched.iloc[0, 11])
                    except Exception: pass

                deliv_sub = deliv_df[(deliv_df["date"] == date_str) & (deliv_df["store"] == store_name)]
                day_delivery_sales = int(pd.to_numeric(deliv_sub["uber"], errors="coerce").fillna(0).sum() +
                                         pd.to_numeric(deliv_sub["rocket_now"], errors="coerce").fillna(0).sum() +
                                         pd.to_numeric(deliv_sub["demaekan"], errors="coerce").fillna(0).sum()) if not deliv_sub.empty else 0
                deliv_other_sales = int(pd.to_numeric(deliv_sub["other"], errors="coerce").fillna(0).sum()) if not deliv_sub.empty else 0

                actual_sales = air_sales + day_delivery_sales + air_other_sales + deliv_other_sales
                total_p += p_sales; total_b += b_sales; total_act += actual_sales

            prev_ratio = f"{(total_act / total_p * 100):.1f}%" if total_p > 0 else "-"
            budget_ratio = f"{(total_act / total_b * 100):.1f}%" if total_b > 0 else "-"
            diff_budget = total_act - total_b
            diff_str = f"+{diff_budget:,}円" if diff_budget > 0 else (f"▲{abs(diff_budget):,}円" if diff_budget < 0 else "0円")

            short_label = get_short_store_key(store_name) + "店"
            report_lines.extend([f"【{short_label}】", f"前年比 : {prev_ratio}", f"予算比 : {budget_ratio}", f"累売上 : {total_act:,}円", f"予乖離 : {diff_str}", ""])

        st.success("🎉 報告用テキストを生成しました！以下の枠内をコピーしてLINEに貼り付けてください。")
        st.code("\n".join(report_lines).strip(), language="text")

else:
    st.title("📊 売上月報一覧")
    st.subheader(f"🏬 {selected_store} {year}年{month}月度")
    st.markdown("---")

    bp_df = load_budget_prev_data()
    deliv_df = load_delivery_data()

    def format_currency(val):
        if pd.isna(val): return "-"
        try:
            num = int(val)
            if num < 0: return f"-¥{abs(num):,}"
            elif num > 0: return f"¥{num:,}"
            else: return "¥0"
        except (ValueError, TypeError): return str(val)

    df_month_total = generate_month_df(selected_store, year, month, None, bp_df, deliv_df)
    df_achieved = df_month_total[df_month_total["実績売上"] != 0]
    max_day = 1 if df_achieved.empty else int(df_achieved["日"].max())

    total_p = pd.to_numeric(df_achieved["前年売上"], errors="coerce").fillna(0).sum()
    total_b = pd.to_numeric(df_achieved["予算売上"], errors="coerce").fillna(0).sum()
    total_act = pd.to_numeric(df_achieved["実績売上"], errors="coerce").fillna(0).sum()
    total_air = pd.to_numeric(df_achieved["店頭売上"], errors="coerce").fillna(0).sum()
    total_deliv = pd.to_numeric(df_achieved["デリバリー売上"], errors="coerce").fillna(0).sum()
    total_other = pd.to_numeric(df_achieved["その他売上"], errors="coerce").fillna(0).sum()
    total_pcust = pd.to_numeric(df_achieved["前年客数"], errors="coerce").fillna(0).sum()
    total_acust = pd.to_numeric(df_achieved["実績客数"], errors="coerce").fillna(0).sum()

    tot_pprice = int(total_p / total_pcust) if total_pcust > 0 else 0
    tot_aprice = int(total_act / total_acust) if total_acust > 0 else 0

    p_ratio_str = f"{(total_act / total_p * 100):.1f}%" if total_p > 0 else "-"
    b_ratio_str = f"{(total_act / total_b * 100):.1f}%" if total_b > 0 else "-"
    cust_ratio_str = f"{(total_acust / total_pcust * 100):.1f}%" if total_pcust > 0 else "-"
    price_ratio_str = f"{(tot_aprice / tot_pprice * 100):.1f}%" if tot_pprice > 0 else "-"

    summary_data = [{
        "日": f"{max_day}日時点", "曜日": "合計", "前年比": p_ratio_str, "予算比": b_ratio_str,
        "前年売上": format_currency(total_p), "予算売上": format_currency(total_b), "実績売上": format_currency(total_act),
        "店頭売上": format_currency(total_air), "デリバリー売上": format_currency(total_deliv), "その他売上": format_currency(total_other),
        "前年客数": int(total_pcust), "実績客数": int(total_acust), "客数前年比": cust_ratio_str,
        "前年単価": format_currency(tot_pprice), "実績単価": format_currency(tot_aprice), "単価前年比": price_ratio_str
    }]
    df_summary = pd.DataFrame(summary_data)

    def render_styled_html_table(df, is_summary=False):
        header_bg, header_color, border_color = ("#f39c12" if is_summary else "#f8f9fa"), ("#ffffff" if is_summary else "#333333"), "#e0e0e0"
        html = f'<div style="overflow-x: auto; margin-bottom: 20px;"><table style="width: 100%; border-collapse: collapse; font-family: sans-serif; font-size: 14px; white-space: nowrap;"><thead><tr style="background-color: {header_bg}; color: {header_color}; text-align: center;">'
        for col in df.columns: html += f'<th style="border: 1px solid {border_color}; padding: 8px; font-weight: bold;">{col}</th>'
        html += "</tr></thead><tbody>"
        for idx, row in df.iterrows():
            bg_color = "#fffbf0" if is_summary else ("#fcfcfc" if idx % 2 == 0 else "#ffffff")
            html += f'<tr style="background-color: {bg_color}; text-align: center;">'
            for col in df.columns:
                val = str(row[col])
                style = f"border: 1px solid {border_color}; padding: 8px;"
                if "%" in val:
                    try:
                        num = float(val.replace("%", "").strip())
                        style += " color: #27ae60; font-weight: bold;" if num >= 100.0 else " color: #c0392b; font-weight: bold;"
                    except ValueError: style += " font-weight: bold;"
                elif "-" in val and any(c.isdigit() for c in val): style += " color: #c0392b; font-weight: bold;"
                elif is_summary or col in ["前年売上", "予算売上", "実績売上", "店頭売上", "デリバリー売上", "その他売上", "前年単価", "実績単価"]: style += " font-weight: bold;"
                html += f'<td style="{style}">{val}</td>'
            html += "</tr>"
        html += "</tbody></table></div>"
        return html

    st.markdown(f"### 【{max_day}日時点 累計サマリー】")
    st.markdown(render_styled_html_table(df_summary, is_summary=True), unsafe_allow_html=True)

    st.markdown("---")
    st.markdown("### 🔍 表示する時間帯を選択してください")
    
    if "active_slot_tab" not in st.session_state: st.session_state["active_slot_tab"] = "合計"
    selected_tab = st.radio("時間帯タブ", ["合計", "早番", "中番", "遅番"], horizontal=True, label_visibility="collapsed")
    st.session_state["active_slot_tab"] = selected_tab

    slot_kw_map = {"合計": None, "早番": "早番", "中番": "中番", "遅番": "遅番"}
    df_current_slot = generate_month_df(selected_store, year, month, slot_kw_map[selected_tab], bp_df, deliv_df)
    df_current_slot_formatted = df_current_slot.copy()
    for col in ["前年売上", "予算売上", "実績売上", "店頭売上", "デリバリー売上", "その他売上", "前年単価", "実績単価"]:
        df_current_slot_formatted[col] = df_current_slot_formatted[col].apply(format_currency)

    st.markdown("### 📅 日別実績一覧")
    st.markdown(render_styled_html_table(df_current_slot_formatted), unsafe_allow_html=True)

    st.markdown("---")
    excel_data = create_multi_sheet_excel(selected_store, year, month, bp_df, deliv_df)
    st.download_button(
        label="📥 全時間帯タブ付きExcelをダウンロード", data=excel_data,
        file_name=f"{selected_store}_{year}年月度_売上月報.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary"
    )
